"""核对修订论文、图源与迁移结构；不连接数据库或运行业务测试。"""
from pathlib import Path
from collections import Counter
import argparse
import hashlib
import json
import re
import xml.etree.ElementTree as ET
from docx import Document
from docx.table import Table
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]


def compact(value):
    """忽略排版换行，保留标识、标点与实际字段值。"""
    return re.sub(r'\s|\u200b', '', str(value))


def verify_schema(data):
    """检查全部来源指纹、最终覆盖量及关键后续迁移。"""
    tables = {item['name']: item for item in data['tables']}
    managed = [item for item in tables.values() if not item.get('external')]
    assert Counter(item['service'] for item in managed) == {'admin': 8, 'training': 15, 'learning': 7}
    assert sum(len(item['columns']) for item in managed) == 464
    assert sum(len(item['indexes']) for item in managed) == 110
    paths = {item['path'] for item in data['sources']}
    actual = {p.relative_to(ROOT).as_posix() for p in ROOT.glob('backend/train-*-service/src/main/resources/db/migration/V*.sql')}
    assert paths == actual and len(paths) == 25
    assert len(data['external_sources']) == 1
    assert data['external_sources'][0]['path'] == 'database/design/sys_address.source.sql'
    for source in data['sources'] + data['external_sources']:
        path = ROOT / source['path']
        assert hashlib.sha256(path.read_bytes()).hexdigest() == source['sha256'], path
        assert not re.search(r'\bFOREIGN\s+KEY\b', path.read_text(encoding='utf-8-sig'), re.I), path
    columns = {name: {column['name']: column for column in item['columns']} for name, item in tables.items()}
    assert columns['sys_user']['vehicle_id']['nullable'] is True
    assert columns['sys_user']['vehicle_id']['source'].endswith('V14__user_vehicle_binding.sql')
    assert next(i for i in tables['sys_user']['indexes'] if i['name'] == 'idx_user_enterprise_vehicle')['columns'] == ['enterprise_id', 'vehicle_id']
    for name in ('train_storage_object', 'train_private_image_upload_session'):
        assert columns[name]['enterprise_id']['nullable'] is True
        assert columns[name]['enterprise_id']['source'].endswith('V6__platform_profile_images.sql')
    active = columns['study_session']['active_user_id']
    assert "'FACE_PENDING'" in active['generated'] and active['source'].endswith('V3__face_check_schema.sql')
    for name in ('attendance_action', 'attendance_sequence', 'attendance_verified_at'):
        assert columns['study_session'][name]['source'].endswith('V4__attendance_face_verification.sql')
    for name in ('attendance_photo_object_id', 'sign_in_photo_object_id', 'sign_out_photo_object_id'):
        assert columns['study_session'][name]['nullable'] is True
        assert columns['study_session'][name]['source'].endswith('V5__learning_record_photos.sql')
    assert columns['face_check_log']['photo_object_id']['nullable'] is True
    assert columns['face_check_task']['attempt_count']['default'] == '0'
    assert any(i['kind'] == 'UNIQUE' and i['columns'] == ['enterprise_id', 'task_id'] for i in tables['exam_record']['indexes'])
    for name, expected in [('sys_user_role', ['user_id', 'role_id']), ('sys_role_permission', ['role_id', 'permission_id'])]:
        assert next(i['columns'] for i in tables[name]['indexes'] if i['kind'] == 'PRIMARY') == expected
    address = tables['sys_address']
    assert address['verified'] is True and address['external'] is True
    assert address['engine'] == 'MyISAM' and address['charset'] == 'utf8mb3'
    assert address['collation'] is None and address['auto_increment_counter'] == 4063
    expected_address = [
        ('id', 'BIGINT UNSIGNED', None), ('level', 'TINYINT UNSIGNED', None),
        ('parent_code', 'VARCHAR(64)', "'0'"), ('area_code', 'VARCHAR(64)', "'0'"),
        ('zip_code', 'VARCHAR(64)', "'0'"), ('city_code', 'VARCHAR(64)', "''"),
        ('name', 'VARCHAR(50)', "''"), ('short_name', 'VARCHAR(50)', "''"),
        ('merger_name', 'VARCHAR(50)', "''"), ('pinyin', 'VARCHAR(30)', "''"),
        ('lng', 'DECIMAL(10,6)', "'0.000000'"), ('lat', 'DECIMAL(10,6)', "'0.000000'")]
    assert [(c['name'], c['type'], c['default']) for c in address['columns']] == expected_address
    assert all(c['nullable'] is False and c['foreign_key'] is False for c in address['columns'])
    assert columns['sys_address']['id']['primary'] is True
    assert columns['sys_address']['id']['auto_increment'] is True
    assert not columns['sys_address']['id']['explicit_default']
    for name in ('parent_code', 'area_code', 'zip_code', 'city_code'):
        assert columns['sys_address'][name]['charset'] == 'utf8mb3'
        assert columns['sys_address'][name]['collation'] == 'utf8mb3_general_ci'
    assert [(i['name'], i['kind'], i['columns'], i.get('method')) for i in address['indexes']] == [
        ('PRIMARY', 'PRIMARY', ['id'], None), ('uk_code', 'UNIQUE', ['area_code'], 'BTREE'),
        ('idx_parent_code', 'INDEX', ['parent_code'], 'BTREE')]
    assert len(tables) == 31 and all(item['verified'] for item in tables.values())
    assert sum(len(item['columns']) for item in tables.values()) == 476
    assert sum(len(item['indexes']) for item in tables.values()) == 113
    for item in managed:
        assert 'ENGINE=InnoDB' in item['tail'] and 'utf8mb4_0900_ai_ci' in item['tail']
        assert all(column['foreign_key'] is False for column in item['columns'])
    return tables


def expected_default(column):
    """独立读取结构源中的默认值类别。"""
    if column.get('unverified'):
        return '待核验'
    if column.get('generated'):
        return '生成列'
    if column['explicit_default']:
        return column['default']
    return '隐式NULL' if column['nullable'] else '无'


def verify_word(path, data, tables):
    """逐单元格核对八张核心表和附录三十一张表及全部索引。"""
    document = Document(path)
    counts = Counter()
    numbers = []
    for paragraph in document.paragraphs:
        match = re.fullmatch(r'表(4-\d+|B-\d+) (\w+)表结构', paragraph.text)
        if not match:
            continue
        number, name = match.groups()
        item = tables[name]
        numbers.append(number)
        counts[name] += 1
        element = paragraph._p.getnext()
        assert element.tag == qn('w:tbl'), (number, '缺少字段表')
        table = Table(element, document._body)
        assert len(table.rows) == len(item['columns']) + 1, name
        assert table.rows[0]._tr.find('.//' + qn('w:tblHeader')) is not None, name
        for row, column in zip(table.rows[1:], item['columns']):
            expected = [column['name'], column['type'],
                        ('PK' if column['primary'] else '—') if item['verified'] else '待核验',
                        '待核验' if column['nullable'] is None else ('是' if column['nullable'] else '否'),
                        expected_default(column), column['comment']
                        + ('；更新时自动记录当前时间' if column.get('on_update') else '')
                        + ('；AUTO_INCREMENT自增' if column.get('auto_increment') else '')]
            assert [compact(c.text) for c in row.cells] == [compact(c) for c in expected], (name, column['name'])
            assert row._tr.find('.//' + qn('w:cantSplit')) is not None
        element = element.getnext()
        while element.tag != qn('w:tbl'):
            element = element.getnext()
        index_table = Table(element, document._body)
        assert len(index_table.rows) == max(1, len(item['indexes'])) + 1
        for row, index in zip(index_table.rows[1:], item['indexes']):
            kind = {'PRIMARY': '主键', 'UNIQUE': '唯一', 'INDEX': '普通'}[index['kind']]
            expected = [index['name'], kind + index.get('method', ''), ','.join(index['columns'])]
            assert [compact(cell.text) for cell in row.cells[:3]] == [compact(value) for value in expected], (name, index['name'])
    assert set(numbers) == {f'4-{i}' for i in range(2, 10)} | {f'B-{i}' for i in range(1, 32)}
    assert counts == Counter({name: 2 if name in data['core_tables'] else 1 for name in tables})
    assert len(document.tables) == 91 and len(document.inline_shapes) == 10
    figures = [re.match(r'^图(\d+-\d+)', p.text)[1] for p in document.paragraphs if p.style.name == 'Caption' and p.text.startswith('图')]
    assert figures == ['3-1', '3-2', '3-3'] + [f'4-{i}' for i in range(1, 8)], figures
    assert len(document.sections) == 4
    text = ''.join(p.text for p in document.paragraphs)
    assert all(value in text for value in ('MyISAM', 'utf8mb3_general_ci', 'AUTO_INCREMENT=4063'))
    assert '待核验' not in text and '真实建表定义尚待补充' not in text
    return document


def verify_preserved(before, document):
    """核对与本轮无关的实现、实验、结论、参考文献和原有附录未被重写。"""
    original = Document(before)
    def stable_part(doc):
        texts = [p.text for p in doc.paragraphs if not p.style.name.lower().startswith('toc')]
        start = next(i for i, text in enumerate(texts) if text.startswith('5 ') and '实现' in text)
        stop = next((i for i, text in enumerate(texts) if text == '附录B 数据表结构字典'), len(texts))
        return texts[start:stop]
    assert stable_part(original) == stable_part(document), '第5章及之后的原有正文发生变化'
    for i in (0, 1, 2):
        assert [[c.text for c in row.cells] for row in original.tables[i].rows] == [[c.text for c in row.cells] for row in document.tables[i].rows], ('原有表', i)


def verify_figures():
    """检查图源 XML、标识唯一性与同源导出文件齐全。"""
    directory = ROOT / 'docs/thesis/figures/advisor-20260930'
    files = sorted(directory.glob('*.drawio'))
    assert len(files) == 10
    for path in files:
        root = ET.parse(path).getroot()
        cells = root.findall('.//mxCell')
        assert len({cell.attrib['id'] for cell in cells}) == len(cells)
        assert path.with_suffix('.png').stat().st_size > 1000
        ET.parse(path.with_suffix('.svg'))


def main():
    """验证已保存的论文；可选核对固定输入备份。"""
    parser = argparse.ArgumentParser()
    parser.add_argument('--docx', type=Path)
    parser.add_argument('--before', type=Path)
    args = parser.parse_args()
    path = args.docx or next((ROOT / 'docs/thesis').glob('李剑*.docx'))
    data = json.loads((ROOT / 'database/design/data-dictionary.json').read_text(encoding='utf-8'))
    tables = verify_schema(data)
    document = verify_word(path, data, tables)
    if args.before:
        verify_preserved(args.before, document)
    verify_figures()
    print(json.dumps({'verified_tables': 31, 'pending_tables': [], 'managed_columns': 464,
                      'managed_indexes': 110, 'total_columns': 476, 'total_indexes': 113,
                      'main_core_tables': 8, 'appendix_entries': 31,
                      'figures': 10, 'docx_sha256': hashlib.sha256(path.read_bytes()).hexdigest()}, ensure_ascii=False))


if __name__ == '__main__':
    main()
